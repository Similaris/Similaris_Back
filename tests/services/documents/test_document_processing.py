from io import BytesIO

import pytest
from docx import Document as DocxBuilder
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.database import Base
from app.models.analysis import AnalysisResult, Batch, Document, Segment
from app.models.references import ReferenceDocument, ReferenceSegment
from app.repositories.analysis import SegmentRepository
from app.services.analysis.hybrid_analysis import HybridAnalysisService
from app.services.documents.document_processing import DocumentProcessingService
from app.services.documents.exceptions import RetryableDocumentProcessingError
from app.services.documents.file_storage import store_document_file
from app.services.analysis.hybrid_contracts import DocumentAnalysis
from app.services.references.search_contracts import (
    ReferenceDocumentMatch,
    ReferenceMatch,
    ReferenceSearchOptions,
    ReferenceSearchResult,
    ReferenceSegmentMatch,
)


class FakeHybridAnalysisService:
    def analyze_document(self, document):
        return DocumentAnalysis(
            document_id=document.id,
            total_segments=0,
            analyzed_segments=0,
            suspicious_segments=0,
            segments=(),
            overall_score=0.0,
            suspicious_segment_percentage=0.0,
            lexical_ms=2,
            semantic_ms=3,
        )


def build_docx(*paragraphs: str) -> bytes:
    document = DocxBuilder()
    for text in paragraphs:
        document.add_paragraph(text)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture()
def session_factory(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    monkeypatch.setattr(
        "app.services.analysis.hybrid_analysis.HybridAnalysisService",
        lambda db: FakeHybridAnalysisService(),
    )
    return sessionmaker(bind=engine, autocommit=False, autoflush=False)


def create_pending_document(
    db, filename: str = "paper.docx", content: bytes | None = None
) -> Document:
    batch = db.query(Batch).first()
    if batch is None:
        batch = Batch(status="pendente")
        db.add(batch)
        db.commit()
        db.refresh(batch)

    if content is None:
        content = build_docx("The student wrote an original academic paper.")

    file_path = store_document_file(content, filename, batch.id)
    document = Document(
        batch_id=batch.id,
        filename=filename,
        file_type="docx",
        file_path=file_path,
        status="pendente",
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def test_process_document_completes_and_persists_segments(session_factory):
    db = session_factory()
    document = create_pending_document(db)

    DocumentProcessingService(db).process_document(document.id)

    db.refresh(document)
    assert document.status == "concluido"
    assert document.error_message is None
    assert document.extraction_ms is not None
    assert document.lexical_ms == 2
    assert document.semantic_ms == 3
    assert float(document.plagiarism_percent) == 0.0
    assert document.started_at is not None
    assert document.finished_at is not None

    segments = db.query(Segment).filter_by(document_id=document.id).all()
    assert len(segments) > 0
    assert segments[0].text_clean

    batch = db.get(Batch, document.batch_id)
    assert batch.status == "concluido"
    assert batch.finished_at is not None
    db.close()


def test_process_document_marks_error_when_extraction_fails(session_factory):
    db = session_factory()
    document = create_pending_document(
        db, filename="corrompido.docx", content=b"nao e um docx"
    )

    DocumentProcessingService(db).process_document(document.id)

    db.refresh(document)
    assert document.status == "erro"
    assert document.error_message
    assert document.finished_at is not None
    assert db.query(Segment).filter_by(document_id=document.id).count() == 0

    # Lote com todos os documentos falhados termina em erro.
    batch = db.get(Batch, document.batch_id)
    assert batch.status == "erro"
    db.close()


def test_process_document_marks_error_when_file_is_missing(session_factory):
    db = session_factory()
    document = create_pending_document(db)
    document.file_path = "caminho/inexistente.docx"
    db.commit()

    DocumentProcessingService(db).process_document(document.id)

    db.refresh(document)
    assert document.status == "erro"
    assert "não encontrado" in document.error_message
    db.close()


def test_process_document_ignores_documents_not_pending(session_factory):
    db = session_factory()
    document = create_pending_document(db)
    document.status = "concluido"
    db.commit()

    DocumentProcessingService(db).process_document(document.id)

    db.refresh(document)
    # Nada muda: mensagens reentregues não reprocessam o documento.
    assert document.started_at is None
    assert db.query(Segment).filter_by(document_id=document.id).count() == 0
    assert db.get(Batch, document.batch_id).status == "concluido"
    db.close()


def test_process_document_runs_hybrid_analysis_and_persists_matches(session_factory):
    db = session_factory()
    document = create_pending_document(db)
    reference_document = ReferenceDocument(
        title="Reference",
        source="pan-pc-11",
        language="en",
        corpus_id="source-document00001.txt",
    )
    reference_segment = ReferenceSegment(
        reference_document=reference_document,
        position=1,
        start_offset=0,
        end_offset=14,
        text_original="Reference text",
        text_clean="reference text",
    )
    db.add(reference_document)
    db.commit()

    match = ReferenceMatch(
        source_document=ReferenceDocumentMatch(
            id=reference_document.id,
            corpus_id=reference_document.corpus_id,
            source=reference_document.source,
            language=reference_document.language,
            title=reference_document.title,
        ),
        source_segment=ReferenceSegmentMatch(
            id=reference_segment.id,
            position=1,
            start_offset=0,
            end_offset=14,
            text_original="Reference text",
        ),
        lexical_score=0.82,
        jaccard_score=0.63,
        semantic_score=0.91,
    )

    class FakeReferenceSearch:
        def search(self, text_original, *, mode):
            return ReferenceSearchResult(
                options=ReferenceSearchOptions(mode=mode),
                index_fingerprint="a" * 64,
                matches=(match,),
                comparison_ms=2.0 if mode == "lexical" else 3.0,
            )

    service = DocumentProcessingService(
        db,
        hybrid_analysis_factory=lambda: HybridAnalysisService(
            reference_search=FakeReferenceSearch(),
            segment_repository=SegmentRepository(db),
        ),
    )
    service.process_document(document.id)

    db.refresh(document)
    stored = db.query(AnalysisResult).filter_by(document_id=document.id).all()
    assert document.status == "concluido"
    assert document.lexical_ms == 2
    assert document.semantic_ms == 3
    assert float(document.plagiarism_percent) > 0
    assert len(stored) == 1
    assert stored[0].reference_segment_id == reference_segment.id
    assert float(stored[0].final_score) == pytest.approx(0.8365)
    assert stored[0].is_suspicious is True
    db.close()


def test_process_document_resumes_a_redelivered_processing_document(session_factory):
    db = session_factory()
    document = create_pending_document(db)
    document.status = "processando"
    db.commit()

    DocumentProcessingService(db).process_document(document.id)

    db.refresh(document)
    assert document.status == "concluido"
    assert db.query(Segment).filter_by(document_id=document.id).count() > 0
    db.close()


def test_transient_analysis_failure_is_retried_then_can_be_finalized(session_factory):
    db = session_factory()
    document = create_pending_document(db)

    class FailingHybridAnalysisService:
        def analyze_document(self, document):
            raise RuntimeError("reference index temporarily unavailable")

    service = DocumentProcessingService(
        db,
        hybrid_analysis_factory=lambda: FailingHybridAnalysisService(),
    )

    with pytest.raises(RetryableDocumentProcessingError):
        service.process_document(document.id)

    db.refresh(document)
    assert document.status == "processando"

    service.mark_failed_after_retries(document.id, "retries exhausted")
    db.refresh(document)
    assert document.status == "erro"
    assert document.error_message == "retries exhausted"
    assert db.get(Batch, document.batch_id).status == "erro"
    db.close()


def test_process_document_ignores_unknown_document(session_factory):
    db = session_factory()

    DocumentProcessingService(db).process_document(999)

    assert db.query(Document).count() == 0
    db.close()


def test_batch_finishes_only_after_last_document(session_factory):
    db = session_factory()
    first = create_pending_document(db, filename="a.docx")
    second = create_pending_document(db, filename="b.docx")
    service = DocumentProcessingService(db)

    service.process_document(first.id)
    batch = db.get(Batch, first.batch_id)
    assert batch.status == "processando"
    assert batch.finished_at is None

    service.process_document(second.id)
    db.refresh(batch)
    assert batch.status == "concluido"
    assert batch.finished_at is not None
    db.close()


def test_batch_with_partial_failure_still_completes(session_factory):
    db = session_factory()
    failed = create_pending_document(
        db, filename="corrompido.docx", content=b"invalido"
    )
    succeeded = create_pending_document(db, filename="valido.docx")
    service = DocumentProcessingService(db)

    service.process_document(failed.id)
    service.process_document(succeeded.id)

    db.refresh(failed)
    db.refresh(succeeded)
    assert failed.status == "erro"
    assert succeeded.status == "concluido"

    batch = db.get(Batch, failed.batch_id)
    assert batch.status == "concluido"
    db.close()
